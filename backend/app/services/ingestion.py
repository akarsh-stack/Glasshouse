import hashlib
import logging
import uuid

from sqlmodel import Session, select

from ..models import Chunk, Document, utcnow

logger = logging.getLogger(__name__)


class UnsupportedDocument(Exception):
    """The upload could not be read. The caller's problem, so a 4xx, not a 500."""


def _extract_pages(file_bytes: bytes, filename: str) -> list[tuple[str, int | None]]:
    """Return [(text, page_number)], one entry per page for paginated formats.

    Paginated formats yield a page number; formats without pagination yield
    `None`, which is what the UI's `p.{n}` badge keys off. Extraction used to
    join every PDF page into one string before chunking, so `page_number` was
    `None` for every chunk in the system and the badge was unreachable.
    """
    name = filename.lower()
    if name.endswith(".pdf"):
        import io
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError

        try:
            reader = PdfReader(io.BytesIO(file_bytes))
            # 1-based: "page 1" is what a reader sees, and the badge exists to be
            # compared against the source document by a human.
            return [(page.extract_text() or "", i) for i, page in enumerate(reader.pages, 1)]
        except (PdfReadError, OSError, ValueError) as e:
            raise UnsupportedDocument(f"Could not read {filename} as a PDF: {e}") from e

    if name.endswith(".docx"):
        import io
        from docx import Document as DocxDocument
        from docx.opc.exceptions import PackageNotFoundError

        try:
            doc = DocxDocument(io.BytesIO(file_bytes))
        except (PackageNotFoundError, OSError, ValueError, KeyError) as e:
            raise UnsupportedDocument(f"Could not read {filename} as a .docx: {e}") from e
        # .docx has no reliable page breaks without rendering, so it is treated
        # as one unpaginated body rather than guessing at boundaries.
        return [("\n".join(p.text for p in doc.paragraphs), None)]

    return [(file_bytes.decode("utf-8", errors="replace"), None)]


class IngestService:
    def __init__(self, chunker, embedding_service, vector_store):
        self.chunker = chunker
        self.embedding_service = embedding_service
        self.vector_store = vector_store

    async def ingest_document(self, file_bytes: bytes, filename: str, db: Session) -> Document:
        if not file_bytes.strip():
            raise UnsupportedDocument(f"{filename} is empty")

        content_hash = hashlib.sha256(file_bytes).hexdigest()
        existing = db.exec(select(Document).where(Document.content_hash == content_hash)).first()
        if existing:
            return existing

        # Extract and chunk *before* inserting the Document row. A file that
        # can't be read should leave no trace; the previous order inserted a row,
        # then marked it "error" and re-raised, leaving debris in the document
        # rail after every bad upload.
        pages = _extract_pages(file_bytes, filename)
        raw_chunks: list[dict] = []
        for text, page_number in pages:
            if not text.strip():
                continue
            for chunk in self.chunker.chunk(text, {"page_number": page_number}):
                # Chunkers index from 0 on every call, so a per-page index would
                # collide across pages. Renumber across the whole document.
                chunk["chunk_index"] = len(raw_chunks)
                raw_chunks.append(chunk)

        if not raw_chunks:
            raise UnsupportedDocument(f"No readable text found in {filename}")

        doc = Document(
            filename=filename,
            content_hash=content_hash,
            status="processing",
            uploaded_at=utcnow(),
        )
        db.add(doc)
        db.commit()
        db.refresh(doc)

        try:
            texts = [c["text"] for c in raw_chunks]
            embeddings = await self.embedding_service.embed_batch(texts)

            chunk_records = []
            chunk_dicts = []
            for c, emb in zip(raw_chunks, embeddings):
                cid = str(uuid.uuid4())
                chunk_records.append(Chunk(
                    id=cid,
                    document_id=doc.id,
                    text=c["text"],
                    chunk_index=c["chunk_index"],
                    page_number=c.get("page_number"),
                ))
                chunk_dicts.append({
                    "id": cid,
                    "document_id": doc.id,
                    "text": c["text"],
                    "chunk_index": c["chunk_index"],
                    "page_number": c.get("page_number"),
                })

            # Rows first, vectors second. The reverse order can leave vectors in
            # Chroma with no row to delete them by, and `delete_by_document`
            # would then never reach them.
            for cr in chunk_records:
                db.add(cr)
            doc.chunk_count = len(chunk_records)
            doc.status = "ready"
            db.add(doc)
            db.commit()
            self.vector_store.upsert(chunk_dicts, embeddings)
            db.refresh(doc)
        except Exception as e:
            logger.error(f"Ingestion failed for {filename}: {e}")
            doc.status = "error"
            db.add(doc)
            db.commit()
            raise

        return doc
